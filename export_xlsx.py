# -*- coding: utf-8 -*-
"""把 groups.json 导出为 xlsx 表格。

用法：
    python export_xlsx.py                       # 读取同目录 groups.json，输出 groups.xlsx
    python export_xlsx.py -o 分组名单.xlsx       # 指定输出文件名
    python export_xlsx.py data/groups.json -o out.xlsx

输出表头固定为：序号 | 学生1 | 学生2

设计原则与 app.py 一致：数据文件缺失/损坏时**明确报错并退出**，
绝不静默当成空数据导出一张空表（那会让人误以为班级没人填报）。
"""

import argparse
import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # 依赖缺失时给出可操作的提示，而不是抛一堆 traceback
    sys.stderr.write(
        '缺少依赖 openpyxl，请先安装：\n'
        '    pip install openpyxl\n'
    )
    raise SystemExit(1)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DATA_FILE = os.path.join(BASE_DIR, 'groups.json')
DEFAULT_OUTPUT_FILE = os.path.join(BASE_DIR, 'groups.xlsx')

HEADERS = ('序号', '学生1', '学生2')
# 各列宽度（字符数），保证中文姓名不会被挤成 ####
COLUMN_WIDTHS = (8, 22, 22)


class ExportError(Exception):
    """数据文件不可用或不合法时抛出，由 main 统一打印并以非 0 退出码结束。"""


def load_groups(path):
    """读取并校验分组数据，返回已按 id 升序排好的列表。

    - 文件不存在 / 无法读取 / JSON 非法 / 顶层不是列表：一律抛 ExportError
    - 单条记录不是对象、缺少姓名：抛 ExportError 并指出第几条，避免导出出残缺名单
    """
    if not os.path.exists(path):
        raise ExportError('找不到数据文件：{}\n（请确认路径，或先启动服务提交几条数据）'.format(path))

    try:
        # 用 utf-8-sig 读取：既能读 app.py 写出的无 BOM 文件，
        # 也能容忍被记事本等编辑器手工编辑后带上 BOM 的文件
        with open(path, 'r', encoding='utf-8-sig') as f:
            data = json.load(f)
    except json.JSONDecodeError as exc:
        # 与 app.py 处理方式一致：损坏文件应人工介入，这里只报告不修改
        raise ExportError('数据文件不是合法的 JSON：{}（第 {} 行第 {} 列）'.format(
            path, exc.lineno, exc.colno)) from exc
    except OSError as exc:
        raise ExportError('数据文件无法读取：{}'.format(exc)) from exc

    if not isinstance(data, list):
        raise ExportError('数据文件格式异常：顶层应为列表，实际为 {}'.format(type(data).__name__))

    groups = []
    for index, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            raise ExportError('第 {} 条记录不是对象，无法导出：{!r}'.format(index, item))

        name1 = str(item.get('name1') or '').strip()
        name2 = str(item.get('name2') or '').strip()
        if not name1 or not name2:
            raise ExportError('第 {} 条记录缺少学生姓名（name1={!r}, name2={!r}）'.format(
                index, item.get('name1'), item.get('name2')))

        raw_id = item.get('id')
        # id 缺失或非数字时用 0 兜底，保证排序不会因为混合类型而报错
        sort_key = raw_id if isinstance(raw_id, (int, float)) else 0
        groups.append({'id': sort_key, 'name1': name1, 'name2': name2})

    # 按原始组号升序；Python 的 sort 稳定，同 id 的记录保持原有先后顺序
    groups.sort(key=lambda group: group['id'])
    return groups


def build_workbook(groups):
    """根据分组列表构建工作簿：首行表头 + 逐行数据，序号从 1 连续编号"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '分组名单'

    header_font = Font(bold=True, size=12)
    header_fill = PatternFill('solid', fgColor='DCE6F1')
    header_align = Alignment(horizontal='center', vertical='center')
    name_align = Alignment(horizontal='left', vertical='center')
    index_align = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='B7B7B7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    # 序号按导出顺序重新从 1 编号，而不是沿用可能跳号/复用的原始 id
    for offset, group in enumerate(groups):
        row = offset + 2
        values = (offset + 1, group['name1'], group['name2'])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = index_align if column == 1 else name_align
            cell.border = border

    for column, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    # 冻结表头，滚动长名单时表头始终可见；未选中新增单元格
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:{}{}'.format(get_column_letter(len(HEADERS)), len(groups) + 1)
    sheet.sheet_view.selection[0].activeCell = 'A2'
    sheet.sheet_view.selection[0].sqref = 'A2'

    return workbook


def save_workbook(workbook, path):
    """保存工作簿，把常见的“文件被 Excel 占用”转成友好提示"""
    directory = os.path.dirname(os.path.abspath(path))
    if directory and not os.path.isdir(directory):
        raise ExportError('输出目录不存在：{}'.format(directory))

    try:
        workbook.save(path)
    except OSError as exc:
        raise ExportError(
            '写入失败：{}\n（若该文件正在 Excel 中打开，请先关闭再重试）'.format(exc)
        ) from exc


def main(argv=None):
    # Windows 控制台默认不是 UTF-8，先切一下避免中文路径/提示乱码
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, OSError):
            pass

    parser = argparse.ArgumentParser(
        description='将 groups.json 导出为 xlsx（表头：序号 | 学生1 | 学生2）',
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        'data_file', nargs='?', default=DEFAULT_DATA_FILE,
        help='数据文件路径（JSON 列表）',
    )
    parser.add_argument(
        '-o', '--output', default=DEFAULT_OUTPUT_FILE,
        help='输出的 xlsx 文件路径',
    )
    args = parser.parse_args(argv)

    try:
        groups = load_groups(args.data_file)
        workbook = build_workbook(groups)
        save_workbook(workbook, args.output)
    except ExportError as exc:
        sys.stderr.write('导出失败：{}\n'.format(exc))
        return 1

    print('导出成功：{}'.format(os.path.abspath(args.output)))
    print('数据来源：{}'.format(os.path.abspath(args.data_file)))
    print('共导出 {} 条分组记录。'.format(len(groups)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
